import re
from pathlib import Path
from typing import Any

from pypdf import PdfReader
from pypdf.errors import PdfReadError

UPLOADS_DIR = Path(__file__).resolve().parents[2] / "uploads"


def _resolve_uploaded_pdf_path(file_name: str | Path) -> Path:
    if file_name is None or str(file_name).strip() == "":
        raise ValueError("No PDF file name provided.")

    base_dir = UPLOADS_DIR.resolve()
    candidate = Path(str(file_name))
    resolved_path = (
        candidate.resolve() if candidate.is_absolute()
        else (base_dir / candidate).resolve()
    )

    try:
        resolved_path.relative_to(base_dir)
    except ValueError as exc:
        raise ValueError(
            "Only files inside the configured uploads directory are allowed."
        ) from exc

    if not resolved_path.exists():
        raise FileNotFoundError("File does not exist in the uploads directory.")

    if not resolved_path.is_file():
        raise ValueError("The requested path is not a file.")

    if resolved_path.suffix.lower() != ".pdf":
        raise ValueError("Only PDF files are allowed.")

    return resolved_path


def extract_pdf_text_from_upload(file_name: str | Path) -> dict[str, Any]:
    result = {
        "status": "failed",
        "error": "",
        "extracted_text": "",
        "page_count": 0,
        "word_count": 0,
        "meaningful_text_found": False,
        "document_title": "Not extractable",
        "document_headings": [],
    }

    try:
        pdf_path = _resolve_uploaded_pdf_path(file_name)
    except (ValueError, FileNotFoundError) as exc:
        result["error"] = str(exc)
        return result

    try:
        reader = PdfReader(str(pdf_path))
    except (PdfReadError, OSError, ValueError) as exc:
        result["error"] = f"Invalid or corrupt PDF: {exc}"
        return result

    if reader.is_encrypted:
        try:
            decrypt_status = reader.decrypt("")
        except Exception:
            decrypt_status = 0

        if decrypt_status == 0:
            result["error"] = (
                "This PDF is encrypted or password-protected and cannot be read "
                "without a password."
            )
            return result

    pages = []
    try:
        page_count = len(reader.pages)
    except Exception as exc:
        result["error"] = f"Unable to read PDF pages: {exc}"
        return result

    heading_pattern = re.compile(
        r"^(?:(?:\d+[\.\d]*|[A-Z]\.)\s+)?(Abstract|Executive Summary|Introduction|"
        r"Problem Statement|Literature Review|Background|Related Work|Objectives|"
        r"System Architecture|Architecture|Methodology|Proposed Methodology|"
        r"Technical Approach|System Design|Implementation|Experimental Results|"
        r"Results and Discussion|Results|Evaluation|Conclusion|Future Scope|References)\b",
        re.IGNORECASE,
    )

    detected_headings: list[str] = []
    first_page_lines: list[str] = []

    for page_index, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            result["error"] = f"Unable to extract text from page {page_index}: {exc}"
            return result

        clean_text = page_text.strip()
        if clean_text:
            lines = [ln.strip() for ln in clean_text.splitlines() if ln.strip()]
            if page_index == 1:
                first_page_lines = lines
            for ln in lines:
                if heading_pattern.match(ln) and len(ln) < 120:
                    if ln not in detected_headings:
                        detected_headings.append(f"Page {page_index}: {ln}")
            pages.append(f"[Page {page_index}]\n{clean_text}")

    extracted_text = "\n\n".join(pages)
    normalized_text = " ".join(extracted_text.split())
    word_count = len(re.findall(r"\b\w+\b", normalized_text))
    meaningful_text_found = bool(re.search(r"[A-Za-z0-9]", normalized_text))

    if not meaningful_text_found:
        result["status"] = "no_text_found"
        result["error"] = (
            "No selectable text found. The PDF may be scanned, image-only, or empty; "
            "OCR is not implemented in this phase."
        )
    else:
        result["status"] = "success"
        result["error"] = ""

    # Determine document title from metadata or first page header
    doc_title = "Not extractable"
    try:
        meta_title = reader.metadata.title if reader.metadata else None
        if meta_title and str(meta_title).strip() and len(str(meta_title).strip()) > 3:
            cand = str(meta_title).strip()
            if not any(g in cand.lower() for g in ("untitled", "microsoft word", "powerpoint")):
                doc_title = cand
    except Exception:
        pass

    if doc_title == "Not extractable" and first_page_lines:
        for candidate in first_page_lines[:4]:
            if len(candidate) > 4 and len(candidate) < 180 and not candidate.lower().startswith("page"):
                doc_title = candidate
                break

    result["extracted_text"] = extracted_text
    result["page_count"] = page_count
    result["word_count"] = word_count
    result["meaningful_text_found"] = meaningful_text_found
    result["document_title"] = doc_title
    result["document_headings"] = detected_headings

    return result


def _resolve_uploaded_presentation_path(file_name: str | Path) -> Path:
    if file_name is None or str(file_name).strip() == "":
        raise ValueError("No presentation file name provided.")

    base_dir = UPLOADS_DIR.resolve()
    candidate = Path(str(file_name))
    resolved_path = (
        candidate.resolve() if candidate.is_absolute()
        else (base_dir / candidate).resolve()
    )

    try:
        resolved_path.relative_to(base_dir)
    except ValueError as exc:
        raise ValueError(
            "Only files inside the configured uploads directory are allowed."
        ) from exc

    if not resolved_path.exists():
        raise FileNotFoundError("File does not exist in the uploads directory.")

    if not resolved_path.is_file():
        raise ValueError("The requested path is not a file.")

    ext = resolved_path.suffix.lower()
    if ext not in {".pptx", ".ppt"}:
        raise ValueError("Only presentation files (.pptx, .ppt) are allowed.")

    return resolved_path


def extract_pptx_text_from_upload(file_name: str | Path) -> dict[str, Any]:
    """Extract slide-by-slide text from an uploaded PPTX file, preserving slide numbers and titles."""
    result = {
        "status": "failed",
        "error": "",
        "extracted_text": "",
        "page_count": 0,
        "word_count": 0,
        "meaningful_text_found": False,
        "document_title": "Not extractable",
        "first_slide_title": "Not extractable",
        "document_headings": [],
    }

    try:
        pres_path = _resolve_uploaded_presentation_path(file_name)
    except (ValueError, FileNotFoundError) as exc:
        result["error"] = str(exc)
        return result

    if pres_path.suffix.lower() == ".ppt":
        result["status"] = "unsupported_format"
        result["error"] = (
            "Legacy PowerPoint (.ppt) format is not supported for text extraction. "
            "Please convert the presentation to .pptx or .pdf format."
        )
        return result

    try:
        from pptx import Presentation
    except ImportError:
        result["error"] = "python-pptx library is not available in the backend environment."
        return result

    try:
        prs = Presentation(str(pres_path))
    except Exception as exc:
        result["error"] = f"Invalid or corrupt PPTX presentation: {exc}"
        return result

    slides_content: list[str] = []
    document_headings: list[str] = []
    first_slide_title = "Not extractable"
    document_title = "Not extractable"
    total_slides = len(prs.slides)

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            cand = slide.shapes.title.text.strip()
            if cand:
                slide_title = cand

        slide_lines: list[str] = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                if not slide_title and shape.text_frame.text.strip():
                    cand = shape.text_frame.text.strip().splitlines()[0].strip()
                    if len(cand) < 140:
                        slide_title = cand
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_lines.append(line)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        slide_lines.append(" | ".join(row_texts))

        if slide_title:
            document_headings.append(f"Slide {slide_idx}: {slide_title}")
        elif slide_lines and len(slide_lines[0]) < 140:
            document_headings.append(f"Slide {slide_idx}: {slide_lines[0]}")

        if slide_idx == 1:
            first_slide_title = slide_title or (slide_lines[0] if slide_lines and len(slide_lines[0]) < 180 else "Not extractable")
            document_title = first_slide_title

        heading_label = f"Slide {slide_idx}: {slide_title}" if slide_title else f"Slide {slide_idx}"
        if slide_lines:
            slides_content.append(f"[{heading_label}]\n" + "\n".join(slide_lines))

    extracted_text = "\n\n".join(slides_content)
    normalized_text = " ".join(extracted_text.split())
    word_count = len(re.findall(r"\b\w+\b", normalized_text))
    meaningful_text_found = bool(re.search(r"[A-Za-z0-9]", normalized_text))

    if not meaningful_text_found:
        result["status"] = "no_text_found"
        result["error"] = (
            "No readable text found in presentation slides. The presentation may contain "
            "only images, shapes without text, or empty slides; OCR is not implemented in this phase."
        )
    else:
        result["status"] = "success"
        result["error"] = ""

    if document_title == "Not extractable":
        try:
            core_t = prs.core_properties.title if hasattr(prs, "core_properties") else None
            if core_t and len(str(core_t).strip()) > 3:
                document_title = str(core_t).strip()
        except Exception:
            pass

    result["extracted_text"] = extracted_text
    result["page_count"] = total_slides
    result["word_count"] = word_count
    result["meaningful_text_found"] = meaningful_text_found
    result["document_title"] = document_title
    result["first_slide_title"] = first_slide_title
    result["document_headings"] = document_headings

    return result


def extract_docx_text_from_upload(file_name: str | Path) -> dict[str, Any]:
    """Extract text from an uploaded DOCX file using standard library zipfile and xml parsing."""
    import xml.etree.ElementTree as ET
    import zipfile

    result = {
        "status": "failed",
        "error": "",
        "extracted_text": "",
        "page_count": 1,
        "word_count": 0,
        "meaningful_text_found": False,
        "document_title": "Not extractable",
        "document_headings": [],
    }

    try:
        base_dir = UPLOADS_DIR.resolve()
        candidate = Path(str(file_name))
        resolved_path = (
            candidate.resolve() if candidate.is_absolute()
            else (base_dir / candidate).resolve()
        )
        resolved_path.relative_to(base_dir)
        if not resolved_path.is_file():
            result["error"] = "File does not exist in uploads directory."
            return result
    except Exception as exc:
        result["error"] = str(exc)
        return result

    try:
        with zipfile.ZipFile(resolved_path) as docx_zip:
            if "word/document.xml" not in docx_zip.namelist():
                result["error"] = "Invalid DOCX format: word/document.xml missing."
                return result
            xml_content = docx_zip.read("word/document.xml")
            tree = ET.fromstring(xml_content)

            paragraphs = []
            headings = []
            for p in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
                texts = [node.text for node in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t") if node.text]
                p_text = "".join(texts).strip()
                if p_text:
                    paragraphs.append(p_text)
                    if len(p_text) < 100 and any(h in p_text.lower() for h in ["abstract", "introduction", "architecture", "requirements", "conclusion", "scope"]):
                        headings.append(p_text)

            full_text = "\n\n".join(paragraphs)
            words = len(re.findall(r"\b\w+\b", full_text))
            doc_title = paragraphs[0] if paragraphs and len(paragraphs[0]) < 120 else "Not extractable"

            result["status"] = "success" if full_text else "no_text_found"
            result["extracted_text"] = full_text
            result["word_count"] = words
            result["meaningful_text_found"] = bool(re.search(r"[A-Za-z0-9]", full_text))
            result["document_title"] = doc_title
            result["document_headings"] = headings
            return result
    except Exception as exc:
        result["error"] = f"Unable to read DOCX document: {exc}"
        return result


def extract_document_text_from_upload(file_name: str | Path) -> dict[str, Any]:
    """Unified document text extractor for all supported file types (.pdf, .pptx, .docx)."""
    if file_name is None or str(file_name).strip() == "":
        return {
            "status": "failed",
            "error": "No file name provided.",
            "extracted_text": "",
            "page_count": 0,
            "word_count": 0,
            "meaningful_text_found": False,
            "document_title": "Not extractable",
            "document_headings": [],
        }

    suffix = Path(str(file_name)).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf_text_from_upload(file_name)
    if suffix in {".pptx", ".ppt"}:
        return extract_pptx_text_from_upload(file_name)
    if suffix == ".docx":
        return extract_docx_text_from_upload(file_name)

    return {
        "status": "unsupported_format",
        "error": f"File extension '{suffix}' is not supported for text extraction. Supported formats are PDF, PPTX, and DOCX.",
        "extracted_text": "",
        "page_count": 0,
        "word_count": 0,
        "meaningful_text_found": False,
        "document_title": "Not extractable",
        "document_headings": [],
    }
